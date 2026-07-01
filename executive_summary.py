from __future__ import annotations

import html
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SEVERITY_BY_ISSUE_TYPE: Dict[str, str] = {
    "DEFERRAL_MISMATCH": "Medium",
    "LOAN_MISMATCH": "Medium",
    "ONLY_IN_PAYROLL": "Medium",
    "ONLY_IN_RECORDKEEPER": "Medium",
    "LATE_CONTRIBUTION": "Low",
    "EMPLOYMENT_STATUS_CONFLICT": "High",
    "POST_TERMINATION_COMPENSATION": "High",
    "HCE catch-up not coded as Roth": "High",
    "402(g) excess deferrals": "High",
    "Under-match from compensation definition variance": "High",
    "Over-match from excluded compensation included": "Low",
    "Match paid to excluded employee class": "High",
    "Under-match": "High",
    "Over-match": "Low",
}


DISPLAY_LABEL_BY_ISSUE_TYPE: Dict[str, str] = {
    'DEFERRAL_MISMATCH': 'Deferral amount mismatch',
    'LOAN_MISMATCH': 'Loan payment mismatch',
    'ONLY_IN_PAYROLL': 'Participant found only in payroll',
    'ONLY_IN_RECORDKEEPER': 'Participant found only in recordkeeper',
    'EMPLOYMENT_STATUS_CONFLICT': 'Employment status conflict',
    'POST_TERMINATION_COMPENSATION': 'Compensation after termination',
    'LATE_CONTRIBUTION': 'Late contribution deposit',
    'SECURE20_HCE_PRETAX_CATCHUP': 'HCE catch-up coded incorrectly',
    'HCE catch-up not coded as Roth': 'HCE catch-up coded incorrectly',
    'COMP_MATCH_UNDER': 'Employer match underpayment',
    'Under-match from compensation definition variance': 'Employer match underpayment',
    'Under-match': 'Employer match underpayment',
    'COMP_MATCH_OVER': 'Employer match overpayment',
    'Over-match from excluded compensation included': 'Employer match overpayment',
    'Over-match': 'Employer match overpayment',
    '402(g) excess deferrals': '402(g) excess deferral',
    'Match paid to excluded employee class': 'Employer match paid to excluded class',
}

WHY_IT_MATTERS: Dict[str, str] = {
    'DEFERRAL_MISMATCH': 'Payroll and recordkeeper totals do not agree for the participant.',
    'LOAN_MISMATCH': 'Loan repayments may be missing, duplicated, or posted to the wrong record.',
    'ONLY_IN_PAYROLL': 'A participant appears in payroll activity but not in the recordkeeper file.',
    'ONLY_IN_RECORDKEEPER': 'A participant appears in the recordkeeper file but not in payroll activity.',
    'LATE_CONTRIBUTION': 'Deposits posted later than the expected remittance window.',
    'EMPLOYMENT_STATUS_CONFLICT': 'Payroll and recordkeeper status values disagree for the same participant.',
    'POST_TERMINATION_COMPENSATION': 'Compensation appears after the termination date held by the recordkeeper.',
    'SECURE20_HCE_PRETAX_CATCHUP': 'HCE catch-up contributions may need Roth source coding under Secure 2.0.',
    'HCE catch-up not coded as Roth': 'HCE catch-up contributions may need Roth source coding under Secure 2.0.',
    'Under-match from compensation definition variance': 'The participant may have received less employer match than the plan formula indicates.',
    'Under-match': 'The participant may have received less employer match than the plan formula indicates.',
    'Over-match from excluded compensation included': 'The participant may have received more employer match than the plan formula indicates.',
    'Over-match': 'The participant may have received more employer match than the plan formula indicates.',
    '402(g) excess deferrals': 'Annual elective deferrals may exceed the statutory limit.',
    'Match paid to excluded employee class': 'Employer match appears for a class that may be excluded under the plan rules.',
}


RECOMMENDED_ACTIONS: Dict[str, str] = {
    "DEFERRAL_MISMATCH": "Reconcile payroll and recordkeeper contribution amounts and correct the affected source or deposit record.",
    "LOAN_MISMATCH": "Review payroll loan withholding against recordkeeper posting and correct any missing or misposted repayment.",
    "ONLY_IN_PAYROLL": "Confirm whether the payroll participant should have been transmitted to the recordkeeper and correct the missing record if needed.",
    "ONLY_IN_RECORDKEEPER": "Confirm whether the recordkeeper-only participant belongs to this plan and payroll cycle, then correct the source feed if needed.",
    "LATE_CONTRIBUTION": "Review payroll remittance timing and determine whether operational corrective action is required.",
    "EMPLOYMENT_STATUS_CONFLICT": "Confirm current participant status and update the incorrect source system or census feed.",
    "POST_TERMINATION_COMPENSATION": "Validate the compensation payment and termination date, then correct payroll or recordkeeper records as needed.",
    "HCE catch-up not coded as Roth": "Confirm HCE status and catch-up source coding, then reclassify catch-up dollars to Roth if required.",
    "402(g) excess deferrals": "Review annual elective deferrals against the statutory limit and coordinate correction of any excess amount.",
    "Under-match from compensation definition variance": "Recalculate employer match under the plan formula and correct any under-credited match amount.",
    "Over-match from excluded compensation included": "Confirm eligible compensation and determine whether over-credited match dollars need operational correction.",
    "Match paid to excluded employee class": "Review employee class eligibility and correct employer match paid to an excluded class if confirmed.",
    "Under-match": "Reconcile employer match calculations against plan formula and correct any under-credited match amount.",
    "Over-match": "Review employer match overage and determine whether payroll or recordkeeper records require adjustment.",
}


ISSUE_CATEGORIES: Dict[str, str] = {
    "DEFERRAL_MISMATCH": "Core Reconciliation",
    "LOAN_MISMATCH": "Core Reconciliation",
    "ONLY_IN_PAYROLL": "Core Reconciliation",
    "ONLY_IN_RECORDKEEPER": "Core Reconciliation",
    "LATE_CONTRIBUTION": "Contribution Timing",
}


ALL_EXCEPTION_COLUMNS = [
    "priority",
    "finding_name",
    "participant",
    "why_it_matters",
    "recommended_action",
    "issue_category",
    "issue_type",
    "employee_id",
    "source",
    "details",
]

def _safe_read_csv(path: Optional[str | Path]) -> pd.DataFrame:
    if not path:
        return pd.DataFrame()
    path = Path(path)
    if not path.exists() or not path.is_file():
        return pd.DataFrame()
    try:
        return pd.read_csv(path)
    except Exception:
        return pd.DataFrame()


def _as_int(value: Any) -> int:
    try:
        if pd.isna(value):
            return 0
        return int(value)
    except Exception:
        return 0


def _severity(issue_type: str, fallback: Any = None) -> str:
    issue_key = str(issue_type)
    if issue_key in SEVERITY_BY_ISSUE_TYPE:
        return SEVERITY_BY_ISSUE_TYPE[issue_key]
    if fallback is not None and not pd.isna(fallback):
        value = str(fallback).strip().title()
        if value in {'High', 'Medium', 'Low', 'Informational'}:
            return value
    return 'Medium'

def _display_label(issue_type: str) -> str:
    return DISPLAY_LABEL_BY_ISSUE_TYPE.get(str(issue_type), str(issue_type).replace('_', ' ').title())


def _why_it_matters(issue_type: str, details: Any = None) -> str:
    issue_key = str(issue_type)
    if issue_key in WHY_IT_MATTERS:
        return WHY_IT_MATTERS[issue_key]
    if details is not None and not pd.isna(details) and str(details).strip():
        return str(details).strip()
    return 'This finding needs operational review before relying on the file for plan administration.'


def _action(issue_type: str, fallback: Any = None) -> str:
    if fallback is not None and not pd.isna(fallback) and str(fallback).strip():
        return str(fallback).strip()
    return RECOMMENDED_ACTIONS.get(
        str(issue_type),
        "Review the supporting detail and coordinate operational follow-up with the responsible source owner.",
    )


def _employee_id(row: pd.Series) -> Optional[str]:
    value = row.get("employee_id")
    if value is None or pd.isna(value):
        return None
    return str(value)


def _append_rows_from_csv(
    rows: List[Dict[str, Any]],
    *,
    path: Optional[str | Path],
    issue_type: str,
    category: str,
    source: str,
    detail_fields: Iterable[str],
    row_filter: Optional[str] = None,
) -> None:
    df = _safe_read_csv(path)
    if df.empty:
        return
    if row_filter and row_filter in df.columns:
        df = df[df[row_filter].astype(str).str.lower().isin({"true", "1"})].copy()
    for _, row in df.iterrows():
        details = []
        for field in detail_fields:
            if field in row.index and not pd.isna(row.get(field)):
                details.append(f"{field}={row.get(field)}")
        detail_text = ", ".join(details)
        participant = _employee_id(row)
        rows.append(
            {
                "priority": _severity(issue_type),
                "finding_name": _display_label(issue_type),
                "participant": participant,
                "why_it_matters": _why_it_matters(issue_type, detail_text),
                "recommended_action": _action(issue_type),
                "issue_category": category,
                "issue_type": issue_type,
                "employee_id": participant,
                "source": source,
                "details": detail_text,
            }
        )


def build_all_exception_rows(
    *,
    output_dir: Path,
    reconciliation_results: Dict[str, Any],
    secure20_summary: Dict[str, Any],
    comp_402g_summary: Dict[str, Any],
    compensation_match_summary: Dict[str, Any],
    population_validation_summary: Dict[str, Any],
    timing_result: Dict[str, Any],
) -> List[Dict[str, Any]]:
    output_dir = Path(output_dir)
    rows: List[Dict[str, Any]] = []

    _append_rows_from_csv(
        rows,
        path=reconciliation_results.get("deferral_mismatches"),
        issue_type="DEFERRAL_MISMATCH",
        category=ISSUE_CATEGORIES["DEFERRAL_MISMATCH"],
        source="deferral_mismatches",
        detail_fields=["pay_date", "amount_payroll", "amount_recordkeeper", "amount_diff"],
    )
    _append_rows_from_csv(
        rows,
        path=reconciliation_results.get("loan_mismatches"),
        issue_type="LOAN_MISMATCH",
        category=ISSUE_CATEGORIES["LOAN_MISMATCH"],
        source="loan_mismatches",
        detail_fields=["pay_date", "amount_payroll", "amount_recordkeeper", "amount_diff"],
    )
    _append_rows_from_csv(
        rows,
        path=reconciliation_results.get("only_in_payroll"),
        issue_type="ONLY_IN_PAYROLL",
        category=ISSUE_CATEGORIES["ONLY_IN_PAYROLL"],
        source="only_in_payroll_deferrals",
        detail_fields=["amount_payroll", "pay_date"],
    )
    _append_rows_from_csv(
        rows,
        path=reconciliation_results.get("only_in_recordkeeper"),
        issue_type="ONLY_IN_RECORDKEEPER",
        category=ISSUE_CATEGORIES["ONLY_IN_RECORDKEEPER"],
        source="only_in_recordkeeper_deferrals",
        detail_fields=["amount_recordkeeper", "pay_date"],
    )
    _append_rows_from_csv(
        rows,
        path=timing_result.get("late_contributions_path") or output_dir / "late_contributions.csv",
        issue_type="LATE_CONTRIBUTION",
        category=ISSUE_CATEGORIES["LATE_CONTRIBUTION"],
        source="late_contributions",
        detail_fields=["pay_date", "deposit_date", "days_to_deposit"],
        row_filter="is_late",
    )

    for source_name, summary in [
        ("secure20", secure20_summary),
        ("comp_402g", comp_402g_summary),
        ("compensation_match", compensation_match_summary),
        ("population_validation", population_validation_summary),
    ]:
        df = _safe_read_csv(summary.get("csv_path") if isinstance(summary, dict) else None)
        if df.empty:
            continue
        for _, row in df.iterrows():
            issue_type = str(
                row.get("exception_type")
                or row.get("violation_type")
                or row.get("issue_type")
                or source_name
            )
            category = str(row.get("issue_category") or ISSUE_CATEGORIES.get(issue_type, source_name))
            details = row.get("details")
            if details is None or pd.isna(details):
                details = ""
            participant = _employee_id(row)
            rows.append(
                {
                    "priority": _severity(issue_type, row.get("severity")),
                    "finding_name": _display_label(issue_type),
                    "participant": participant,
                    "why_it_matters": _why_it_matters(issue_type, details),
                    "recommended_action": _action(issue_type, row.get("correction_hint")),
                    "issue_category": category,
                    "issue_type": issue_type,
                    "employee_id": participant,
                    "source": source_name,
                    "details": str(details),
                }
            )

    return rows


def _review_counts(df: pd.DataFrame) -> Tuple[int, int]:
    if df is None or df.empty:
        return 0, 0
    employees = int(df["employee_id"].astype(str).str.strip().nunique()) if "employee_id" in df.columns else 0
    periods = 0
    if "pay_period_number" in df.columns:
        periods = int(df["pay_period_number"].nunique())
    elif "pay_date" in df.columns:
        periods = int(pd.to_datetime(df["pay_date"], errors="coerce").dropna().dt.normalize().nunique())
    return employees, periods


def _collect_warnings(*summaries: Dict[str, Any]) -> List[str]:
    warnings: List[str] = []
    for summary in summaries:
        if not isinstance(summary, dict):
            continue
        warning = summary.get("warning")
        if warning:
            warnings.append(str(warning))
        skipped = summary.get("skipped_rules")
        if isinstance(skipped, dict):
            for key, reason in skipped.items():
                warnings.append(f"{key}: {reason}")
    return warnings


def build_executive_summary(
    *,
    run_id: str,
    plan_name: Optional[str],
    plan_year: Optional[int],
    payroll_df: pd.DataFrame,
    recordkeeper_df: pd.DataFrame,
    output_dir: Path,
    reconciliation_results: Dict[str, Any],
    secure20_summary: Dict[str, Any],
    eligibility_summary: Dict[str, Any],
    comp_402g_summary: Dict[str, Any],
    match_summary: Dict[str, Any],
    compensation_match_summary: Dict[str, Any],
    population_validation_summary: Dict[str, Any],
    timing_result: Dict[str, Any],
    evidence_pack_path: Optional[str],
    verification_status: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    employees_reviewed, payroll_periods_reviewed = _review_counts(payroll_df)
    exception_rows = build_all_exception_rows(
        output_dir=Path(output_dir),
        reconciliation_results=reconciliation_results,
        secure20_summary=secure20_summary,
        comp_402g_summary=comp_402g_summary,
        compensation_match_summary=compensation_match_summary,
        population_validation_summary=population_validation_summary,
        timing_result=timing_result,
    )
    exceptions_df = pd.DataFrame(exception_rows, columns=ALL_EXCEPTION_COLUMNS)
    if exceptions_df.empty:
        exceptions_df = pd.DataFrame(columns=ALL_EXCEPTION_COLUMNS)

    priority_counts = exceptions_df["priority"].value_counts().to_dict() if not exceptions_df.empty else {}
    category_counts = exceptions_df["issue_category"].value_counts().to_dict() if not exceptions_df.empty else {}
    type_counts = exceptions_df["finding_name"].value_counts().to_dict() if not exceptions_df.empty else {}

    priority_rank = {"High": 0, "Medium": 1, "Low": 2, "Informational": 3}
    top_df = exceptions_df.copy()
    if not top_df.empty:
        top_df["_rank"] = top_df["priority"].map(priority_rank).fillna(3)
        top_df = top_df.sort_values(["_rank", "issue_category", "finding_name"]).drop(columns=["_rank"])

    recommended_actions = []
    seen_actions = set()
    for _, row in top_df.iterrows():
        action = str(row.get("recommended_action", "")).strip()
        if action and action not in seen_actions:
            recommended_actions.append(
                {
                    "finding_name": row.get("finding_name"),
                    "issue_type": row.get("issue_type"),
                    "priority": row.get("priority"),
                    "action": action,
                }
            )
            seen_actions.add(action)

    verification = verification_status or {
        "input_verification": "PENDING",
        "output_verification": "PENDING",
        "overall_verification": "PENDING",
        "message": "Evidence pack generated; run verify_proof.py for independent verification.",
    }

    return {
        "run_id": run_id,
        "plan_name": plan_name or "Unknown Plan",
        "plan_year": plan_year,
        "generated_timestamp": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "employees_reviewed": int(employees_reviewed),
        "payroll_periods_reviewed": int(payroll_periods_reviewed),
        "payroll_row_count": int(len(payroll_df)) if payroll_df is not None else 0,
        "recordkeeper_row_count": int(len(recordkeeper_df)) if recordkeeper_df is not None else 0,
        "total_exceptions": int(len(exceptions_df)),
        "high_priority_count": int(priority_counts.get("High", 0)),
        "medium_priority_count": int(priority_counts.get("Medium", 0)),
        "low_priority_count": int(priority_counts.get("Low", 0)),
        "counts_by_issue_category": {str(k): int(v) for k, v in category_counts.items()},
        "counts_by_issue_type": {str(k): int(v) for k, v in type_counts.items()},
        "top_priority_exceptions": top_df.head(10).to_dict("records"),
        "recommended_actions": recommended_actions[:10],
        "skipped_analyzer_warnings": _collect_warnings(
            secure20_summary,
            eligibility_summary,
            comp_402g_summary,
            match_summary,
            compensation_match_summary,
            population_validation_summary,
        ),
        "evidence_pack_path": evidence_pack_path,
        "verification_status": verification,
        "synthetic_data_disclosure": "All data is synthetic." if "polished" in str(evidence_pack_path).lower() else None,
    }


def write_summary_artifacts(
    summary: Dict[str, Any],
    exception_rows: List[Dict[str, Any]],
    output_dir: Path,
) -> Dict[str, str]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "executive_summary.json"
    html_path = output_dir / "executive_summary.html"
    csv_path = output_dir / "all_exceptions.csv"

    pd.DataFrame(exception_rows, columns=ALL_EXCEPTION_COLUMNS).to_csv(csv_path, index=False)

    temp_json = json_path.with_suffix(".json.tmp")
    temp_json.write_text(json.dumps(summary, indent=2, default=str), encoding="utf-8")
    temp_json.replace(json_path)

    html_text = render_summary_html(summary)
    temp_html = html_path.with_suffix(".html.tmp")
    temp_html.write_text(html_text, encoding="utf-8")
    temp_html.replace(html_path)

    return {
        "executive_summary_json": str(json_path),
        "executive_summary_html": str(html_path),
        "all_exceptions_csv": str(csv_path),
    }


def render_summary_html(summary: Dict[str, Any]) -> str:
    def esc(value: Any) -> str:
        return html.escape("" if value is None else str(value))

    categories = "".join(
        f"<tr><td>{esc(k)}</td><td>{int(v)}</td></tr>"
        for k, v in summary.get("counts_by_issue_category", {}).items()
    )
    findings = "".join(
        "<tr>"
        f"<td>{esc(item.get('priority'))}</td>"
        f"<td>{esc(item.get('issue_category'))}</td>"
        f"<td>{esc(item.get('finding_name'))}</td>"
        f"<td>{esc(item.get('participant'))}</td>"
        f"<td>{esc(item.get('why_it_matters'))}</td>"
        f"<td>{esc(item.get('recommended_action'))}</td>"
        "</tr>"
        for item in summary.get("top_priority_exceptions", [])
    )
    actions = "".join(
        f"<li><strong>{esc(item.get('finding_name'))}:</strong> {esc(item.get('action'))}</li>"
        for item in summary.get("recommended_actions", [])
    )
    verification = summary.get("verification_status") or {}
    disclosure = summary.get("synthetic_data_disclosure")
    disclosure_html = f"<p class='note'>{esc(disclosure)}</p>" if disclosure else ""

    return f"""<!doctype html>
<html>
<head>
  <meta charset="utf-8">
  <title>ProofLink Executive Summary</title>
  <style>
    body {{ font-family: Arial, sans-serif; margin: 32px; color: #1f2933; }}
    h1 {{ margin-bottom: 4px; }}
    h2 {{ margin-top: 28px; border-bottom: 1px solid #d8dee9; padding-bottom: 4px; }}
    .metrics {{ display: grid; grid-template-columns: repeat(4, 1fr); gap: 12px; }}
    .metric {{ border: 1px solid #d8dee9; padding: 12px; border-radius: 6px; }}
    .label {{ color: #52616b; font-size: 12px; text-transform: uppercase; }}
    .value {{ font-size: 24px; font-weight: 700; }}
    table {{ width: 100%; border-collapse: collapse; margin-top: 8px; }}
    th, td {{ border: 1px solid #d8dee9; padding: 7px; text-align: left; vertical-align: top; white-space: normal; overflow-wrap: anywhere; }}
    th {{ background: #f3f6f8; }}
    .note {{ color: #52616b; }}
  </style>
</head>
<body>
  <h1>ProofLink Review Summary</h1>
  <p class="note">{esc(summary.get('plan_name'))} | Plan year {esc(summary.get('plan_year'))} | Run {esc(summary.get('run_id'))}</p>
  {disclosure_html}
  <div class="metrics">
    <div class="metric"><div class="label">Employees</div><div class="value">{summary.get('employees_reviewed', 0):,}</div></div>
    <div class="metric"><div class="label">Payroll Periods</div><div class="value">{summary.get('payroll_periods_reviewed', 0):,}</div></div>
    <div class="metric"><div class="label">Exceptions</div><div class="value">{summary.get('total_exceptions', 0):,}</div></div>
    <div class="metric"><div class="label">High Priority</div><div class="value">{summary.get('high_priority_count', 0):,}</div></div>
  </div>
  <h2>Issue Categories</h2>
  <table><thead><tr><th>Category</th><th>Count</th></tr></thead><tbody>{categories}</tbody></table>
  <h2>Top Priority Findings</h2>
  <table><thead><tr><th>Priority</th><th>Category</th><th>Finding</th><th>Participant</th><th>Why it matters</th><th>Recommended action</th></tr></thead><tbody>{findings}</tbody></table>
  <h2>Recommended Actions</h2>
  <ul>{actions}</ul>
  <h2>Evidence Verification</h2>
  <p>Overall verification: <strong>{esc(verification.get('overall_verification'))}</strong></p>
  <p class="note">Detailed workbook tabs and CSV files contain supporting records.</p>
</body>
</html>
"""


def add_review_summary_sheet(excel_path: Path, summary: Dict[str, Any]) -> None:
    excel_path = Path(excel_path)
    if not excel_path.exists():
        return
    wb = load_workbook(excel_path)
    if "Review Summary" in wb.sheetnames:
        del wb["Review Summary"]
    ws = wb.create_sheet("Review Summary", 0)

    title_fill = PatternFill("solid", fgColor="1F4E78")
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    white = Font(color="FFFFFF", bold=True, size=14)
    bold = Font(bold=True)

    ws["A1"] = "ProofLink Review Summary"
    ws["A1"].font = white
    ws["A1"].fill = title_fill
    ws.merge_cells("A1:F1")

    rows = [
        ("Plan name", summary.get("plan_name"), "Plan year", summary.get("plan_year")),
        ("Run ID", summary.get("run_id"), "Generated", summary.get("generated_timestamp")),
        ("Employees reviewed", summary.get("employees_reviewed"), "Payroll periods reviewed", summary.get("payroll_periods_reviewed")),
        ("Payroll rows", summary.get("payroll_row_count"), "Recordkeeper rows", summary.get("recordkeeper_row_count")),
        ("Total exceptions", summary.get("total_exceptions"), "High priority", summary.get("high_priority_count")),
        ("Medium priority", summary.get("medium_priority_count"), "Low priority", summary.get("low_priority_count")),
        ("Verification", (summary.get("verification_status") or {}).get("overall_verification"), "", ""),
    ]
    row_num = 3
    for label1, value1, label2, value2 in rows:
        ws.cell(row_num, 1, label1).font = bold
        ws.cell(row_num, 2, value1)
        ws.cell(row_num, 4, label2).font = bold
        ws.cell(row_num, 5, value2)
        row_num += 1

    row_num += 1
    ws.cell(row_num, 1, "Issue Counts By Category").font = bold
    ws.cell(row_num, 1).fill = section_fill
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=3)
    row_num += 1
    ws.cell(row_num, 1, "Category").font = bold
    ws.cell(row_num, 2, "Count").font = bold
    for category, count in summary.get("counts_by_issue_category", {}).items():
        row_num += 1
        ws.cell(row_num, 1, category)
        ws.cell(row_num, 2, count)

    row_num += 2
    ws.cell(row_num, 1, "Top Priority Exceptions").font = bold
    ws.cell(row_num, 1).fill = section_fill
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=5)
    row_num += 1
    headers = ["Priority", "Finding", "Participant", "Why it matters", "Recommended action"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row_num, col, header).font = bold
    table_start = row_num
    for item in summary.get("top_priority_exceptions", [])[:10]:
        row_num += 1
        ws.cell(row_num, 1, item.get("priority"))
        ws.cell(row_num, 2, item.get("finding_name"))
        ws.cell(row_num, 3, item.get("participant"))
        ws.cell(row_num, 4, item.get("why_it_matters"))
        ws.cell(row_num, 5, item.get("recommended_action"))
    if row_num > table_start:
        ws.auto_filter.ref = f"A{table_start}:E{row_num}"

    row_num += 2
    ws.cell(row_num, 1, "Recommended Next Actions").font = bold
    ws.cell(row_num, 1).fill = section_fill
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=5)
    for action in summary.get("recommended_actions", [])[:8]:
        row_num += 1
        ws.cell(row_num, 1, action.get("finding_name"))
        ws.cell(row_num, 2, action.get("action"))
        ws.merge_cells(start_row=row_num, start_column=2, end_row=row_num, end_column=5)

    row_num += 2
    ws.cell(row_num, 1, "Note").font = bold
    ws.cell(row_num, 2, "Detailed tabs contain supporting records for each output.")
    ws.merge_cells(start_row=row_num, start_column=2, end_row=row_num, end_column=5)

    ws.freeze_panes = "A3"
    widths = {1: 18, 2: 28, 3: 18, 4: 46, 5: 64}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(excel_path)

