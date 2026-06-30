from pathlib import Path
import json
import shutil
import zipfile

import pandas as pd
from openpyxl import load_workbook

import main
from main import EngineConfig, run_prooflink_engine, run_reconciliation
import verify_proof


def _write_csv(path: Path, rows: list[dict]) -> None:
    pd.DataFrame(rows).to_csv(path, index=False)


def _single_manifest_path(proofs_dir: Path) -> Path:
    manifests = list(proofs_dir.glob("proof_manifest_*.json"))
    assert len(manifests) == 1
    return manifests[0]


def test_no_module_level_recordkeeper_dataframe_exists():
    assert not hasattr(main, "_LAST_RECORDKEEPER_PROCESSED_DF")


def test_population_validation_runs_once_per_engine_run(tmp_path: Path, monkeypatch):
    calls = {"count": 0}
    original = main.analyze_population_validation

    def counting_analyzer(*args, **kwargs):
        calls["count"] += 1
        return original(*args, **kwargs)

    monkeypatch.setattr(main, "analyze_population_validation", counting_analyzer)

    payroll_csv = tmp_path / "payroll_once.csv"
    rk_csv = tmp_path / "rk_once.csv"
    _write_csv(
        payroll_csv,
        [
            {
                "employee_id": "9001",
                "pay_date": "2025-03-05",
                "def_amount": 100,
                "roth_amount": 0,
                "loan_amount": 0,
                "employment_status": "active",
                "Compensation": 1200,
            }
        ],
    )
    _write_csv(
        rk_csv,
        [
            {
                "employee_id": "9001",
                "record_date": "2025-01-16",
                "def_amount": 100,
                "roth_amount": 0,
                "loan_amount": 0,
                "employment_status": "terminated",
                "termination_date": "2025-01-31",
            }
        ],
    )

    run_prooflink_engine(
        payroll_path=str(payroll_csv),
        rk_path=str(rk_csv),
        config=EngineConfig(
            plan_name="Population Once Plan",
            payroll_vendor_hint="ADP",
            rk_vendor_hint="VENDOR_RK_1",
            output_dir=str(tmp_path / "output_once"),
            proofs_dir=str(tmp_path / "proofs_once"),
        ),
        run_id="population-validation-once-test",
    )

    assert calls["count"] == 1


def test_sequential_population_runs_are_isolated(tmp_path: Path):
    observed_ids = []
    for employee_id in ["8101", "8202"]:
        run_dir = tmp_path / employee_id
        payroll_csv = run_dir / "payroll.csv"
        rk_csv = run_dir / "rk.csv"
        run_dir.mkdir()
        _write_csv(
            payroll_csv,
            [
                {
                    "employee_id": employee_id,
                    "pay_date": "2025-03-05",
                    "def_amount": 100,
                    "roth_amount": 0,
                    "loan_amount": 0,
                    "employment_status": "active",
                    "Compensation": 1200,
                }
            ],
        )
        _write_csv(
            rk_csv,
            [
                {
                    "employee_id": employee_id,
                    "record_date": "2025-01-16",
                    "def_amount": 100,
                    "roth_amount": 0,
                    "loan_amount": 0,
                    "employment_status": "terminated",
                    "termination_date": "2025-01-31",
                }
            ],
        )
        result = run_prooflink_engine(
            payroll_path=str(payroll_csv),
            rk_path=str(rk_csv),
            config=EngineConfig(
                plan_name=f"Population Isolation {employee_id}",
                payroll_vendor_hint="ADP",
                rk_vendor_hint="VENDOR_RK_1",
                output_dir=str(run_dir / "output"),
                proofs_dir=str(run_dir / "proofs"),
            ),
            run_id=f"population-validation-isolation-{employee_id}",
        )
        issues = pd.read_csv(result.summary["population_validation"]["csv_path"])
        assert set(issues["employee_id"].astype(str)) == {employee_id}
        observed_ids.append(issues.loc[0, "employee_id"])

    assert [str(value) for value in observed_ids] == ["8101", "8202"]


def test_population_validation_engine_outputs_corrected_exceptions(tmp_path: Path):
    payroll_csv = tmp_path / "payroll.csv"
    rk_csv = tmp_path / "rk.csv"
    out_dir = tmp_path / "output"
    proofs_dir = tmp_path / "proofs"

    _write_csv(
        payroll_csv,
        [
            {
                "employee_id": "1001",
                "pay_date": "2025-03-05",
                "def_amount": 100,
                "roth_amount": 0,
                "loan_amount": 0,
                "employment_status": "active",
                "Compensation": 1200,
            },
            {
                "employee_id": "1002",
                "pay_date": "2025-01-15",
                "def_amount": 200,
                "roth_amount": 0,
                "loan_amount": 0,
                "employment_status": "active",
                "Compensation": 1000,
            },
        ],
    )
    _write_csv(
        rk_csv,
        [
            {
                "employee_id": "1001",
                "record_date": "2025-01-16",
                "def_amount": 100,
                "roth_amount": 0,
                "loan_amount": 0,
                "employment_status": "terminated",
                "termination_date": "2025-01-31",
            },
            {
                "employee_id": "1002",
                "record_date": "2025-01-16",
                "def_amount": 200,
                "roth_amount": 0,
                "loan_amount": 0,
                "employment_status": "active",
                "termination_date": "",
            },
        ],
    )

    result = run_prooflink_engine(
        payroll_path=str(payroll_csv),
        rk_path=str(rk_csv),
        config=EngineConfig(
            plan_name="Population Demo Plan",
            payroll_vendor_hint="ADP",
            rk_vendor_hint="VENDOR_RK_1",
            output_dir=str(out_dir),
            proofs_dir=str(proofs_dir),
        ),
        run_id="population-validation-test",
    )

    population = result.summary["population_validation"]
    assert population["issue_count"] == 2
    assert population["employment_status_conflict_count"] == 1
    assert population["post_termination_compensation_count"] == 1
    assert "recordkeeper_processed_df" not in result.summary

    csv_path = Path(population["csv_path"])
    assert csv_path.exists()
    issues = pd.read_csv(csv_path)
    assert set(issues["exception_type"]) == {"EMPLOYMENT_STATUS_CONFLICT", "POST_TERMINATION_COMPENSATION"}
    assert set(issues["issue_category"]) == {"Population Validation"}
    assert set(issues["authoritative_source"]) == {"undetermined"}
    assert set(issues["suspected_origin"]) == {"undetermined"}
    assert set(issues["run_id"]) == {"population-validation-test"}
    assert set(issues["plan_name"]) == {"Population Demo Plan"}
    assert set(issues["plan_year"]) == {2025}

    assert result.summary["plan_exceptions"]["by_category"]["Population Validation"] == 2
    assert any(item.get("key") == "population_validation_issues" for item in result.summary.get("evidence_index", []))

    report_path = out_dir / "reconciliation_report.xlsx"
    workbook = load_workbook(report_path, read_only=True)
    try:
        assert "Population Issues" in workbook.sheetnames
    finally:
        workbook.close()

    manifest_path = _single_manifest_path(proofs_dir)
    with manifest_path.open("r", encoding="utf-8") as f:
        manifest = json.load(f)

    pop_entry = manifest["outputs"]["population_validation_issues"]
    assert not Path(pop_entry["path"]).is_absolute()
    assert Path(pop_entry["path"]).name == csv_path.name
    assert pop_entry["sha256"] == verify_proof.sha256_file(csv_path)
    assert pop_entry["row_count"] == 2
    assert pop_entry["merkle_root"]
    assert pop_entry["row_hash_sample"]

    excel_entry = manifest["outputs"]["excel_report"]
    assert not Path(excel_entry["path"]).is_absolute()
    assert Path(excel_entry["path"]).name == report_path.name
    assert excel_entry["sha256"] == verify_proof.sha256_file(report_path)
    manifest_verification = verify_proof.verify_manifest(manifest_path)
    assert manifest_verification["overall_verification"] == "PASS"
    assert manifest_verification["output_verification"] == "PASS"

    with zipfile.ZipFile(result.evidence_pack_path) as evidence_zip:
        names = set(evidence_zip.namelist())
        assert "population_validation_issues.csv" in names
        assert "plan_exception_summary.csv" in names
        assert manifest_path.name in names
        zipped_manifest = json.loads(evidence_zip.read(manifest_path.name).decode("utf-8"))
        assert zipped_manifest["outputs"]["population_validation_issues"]["path"] == pop_entry["path"]
        assert zipped_manifest["outputs"]["population_validation_issues"]["sha256"] == pop_entry["sha256"]

    extracted_dir = tmp_path / "extracted_pack"
    with zipfile.ZipFile(result.evidence_pack_path) as evidence_zip:
        evidence_zip.extractall(extracted_dir)
    shutil.copy2(payroll_csv, extracted_dir / payroll_csv.name)
    shutil.copy2(rk_csv, extracted_dir / rk_csv.name)
    shutil.rmtree(out_dir)
    payroll_csv.unlink()
    rk_csv.unlink()

    extracted_manifest = extracted_dir / manifest_path.name
    extracted_verification = verify_proof.verify_manifest(extracted_manifest)
    assert extracted_verification["overall_verification"] == "PASS"
    assert extracted_verification["output_verification"] == "PASS"


def test_no_population_validation_issues_manifest_remains_valid(tmp_path: Path):
    payroll_csv = tmp_path / "payroll_clean.csv"
    rk_csv = tmp_path / "rk_clean.csv"
    out_dir = tmp_path / "output_clean"
    proofs_dir = tmp_path / "proofs_clean"

    _write_csv(
        payroll_csv,
        [
            {
                "employee_id": "1001",
                "pay_date": "2025-01-15",
                "def_amount": 100,
                "roth_amount": 0,
                "loan_amount": 0,
                "employment_status": "active",
                "Compensation": 1000,
            }
        ],
    )
    _write_csv(
        rk_csv,
        [
            {
                "employee_id": "1001",
                "record_date": "2025-01-16",
                "def_amount": 100,
                "roth_amount": 0,
                "loan_amount": 0,
                "employment_status": "active",
                "termination_date": "",
            }
        ],
    )

    result = run_prooflink_engine(
        payroll_path=str(payroll_csv),
        rk_path=str(rk_csv),
        config=EngineConfig(
            plan_name="Population Clean Plan",
            payroll_vendor_hint="ADP",
            rk_vendor_hint="VENDOR_RK_1",
            output_dir=str(out_dir),
            proofs_dir=str(proofs_dir),
        ),
        run_id="population-validation-clean-test",
    )

    assert result.summary["population_validation"]["issue_count"] == 0
    assert result.summary["population_validation"]["csv_path"] is None

    manifest_path = _single_manifest_path(proofs_dir)
    with manifest_path.open("r", encoding="utf-8") as f:
        manifest = json.load(f)

    assert "population_validation_issues" not in manifest["outputs"]
    assert Path(result.evidence_pack_path).exists()
    manifest_verification = verify_proof.verify_manifest(manifest_path)
    assert manifest_verification["overall_verification"] == "PASS"
    assert manifest_verification["output_verification"] == "PASS"


def test_existing_payroll_only_and_rk_only_reconciliation_outputs_remain(tmp_path: Path):
    payroll_csv = tmp_path / "payroll_only.csv"
    rk_csv = tmp_path / "rk_only.csv"
    out_dir = tmp_path / "output_recon"
    proofs_dir = tmp_path / "proofs_recon"

    _write_csv(
        payroll_csv,
        [
            {"employee_id": "1001", "pay_date": "2025-01-15", "def_amount": 100, "roth_amount": 0, "loan_amount": 0},
            {"employee_id": "1002", "pay_date": "2025-01-15", "def_amount": 200, "roth_amount": 0, "loan_amount": 0},
        ],
    )
    _write_csv(
        rk_csv,
        [
            {"employee_id": "1001", "record_date": "2025-01-16", "def_amount": 100, "roth_amount": 0, "loan_amount": 0},
            {"employee_id": "1003", "record_date": "2025-01-16", "def_amount": 50, "roth_amount": 0, "loan_amount": 0},
        ],
    )

    results = run_reconciliation(
        payroll_csv=str(payroll_csv),
        rk_csv=str(rk_csv),
        payroll_vendor_hint="ADP",
        rk_vendor_hint="VENDOR_RK_1",
        output_dir=str(out_dir),
        proofs_dir=str(proofs_dir),
        mapping_yaml_path=str(Path(__file__).resolve().parents[1] / "mapping_example.yaml"),
    )

    assert "recordkeeper_processed_df" not in results

    only_payroll = pd.read_csv(results["only_in_payroll"])
    only_rk = pd.read_csv(results["only_in_recordkeeper"])
    assert set(only_payroll["employee_id"].astype(str)) == {"1002"}
    assert set(only_rk["employee_id"].astype(str)) == {"1003"}
