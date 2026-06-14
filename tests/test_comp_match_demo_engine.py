from pathlib import Path
import zipfile

import pandas as pd
from openpyxl import load_workbook

from main import EngineConfig, run_prooflink_engine


ROOT_DIR = Path(__file__).resolve().parents[1]
DEMO_DIR = ROOT_DIR / "demo"
PAYROLL_FILE = DEMO_DIR / "demo_comp_match_payroll.csv"
RK_FILE = DEMO_DIR / "demo_comp_match_rk.csv"
MAPPING_FILE = DEMO_DIR / "demo_comp_match_mapping.yaml"


def test_comp_match_demo_engine_outputs(tmp_path, monkeypatch):
    monkeypatch.setenv("MAPPING_YAML_PATH", str(MAPPING_FILE))

    result = run_prooflink_engine(
        payroll_path=str(PAYROLL_FILE),
        rk_path=str(RK_FILE),
        config=EngineConfig(
            plan_name="Comp Match Demo Plan",
            output_dir=str(tmp_path / "output"),
            proofs_dir=str(tmp_path / "proofs"),
        ),
        run_id="comp-match-demo-test",
        plan_rules={
            "plan_match_config": {
                "match_formula_name": "50% up to 6%",
                "match_type": "percent_of_comp",
                "match_rate": 0.50,
                "match_cap_pct": 0.06,
                "match_frequency": "per_payroll",
                "true_up_enabled": False,
                "true_up_enabled_column": "True-Up Enabled",
                "eligible_comp_columns": ["Regular Compensation", "Overtime", "Bonus"],
                "excluded_comp_columns": ["Fringe", "Reimbursement"],
                "employee_class_column": "Employee Class",
                "eligible_classes": ["Full-Time", "Part-Time"],
                "excluded_classes": ["Intern", "Union Excluded"],
                "absolute_tolerance": 5.00,
                "relative_tolerance_pct": 0.15,
            }
        },
    )

    comp_match = result.summary["compensation_match"]
    assert comp_match["issue_count"] > 0
    assert comp_match["estimated_under_match_dollars"] > 0
    assert comp_match["estimated_over_match_dollars"] > 0
    assert comp_match["excluded_class_match_count"] > 0

    csv_path = Path(comp_match["csv_path"])
    assert csv_path.exists()

    issues = pd.read_csv(csv_path)
    assert "Under-match from compensation definition variance" in set(issues["issue_type"])
    assert "Over-match from excluded compensation included" in set(issues["issue_type"])
    assert "Match paid to excluded employee class" in set(issues["issue_type"])
    assert "Potential annual true-up timing difference" in set(issues["issue_type"])

    true_up_rows = issues[issues["issue_type"] == "Potential annual true-up timing difference"]
    assert not true_up_rows.empty
    assert comp_match["estimated_under_match_dollars"] < issues.loc[
        issues["match_variance"] < 0,
        "match_variance_abs",
    ].sum()

    assert any(
        item.get("key") == "compensation_match_issues"
        for item in result.summary.get("evidence_index", [])
    )

    report_path = tmp_path / "output" / "reconciliation_report.xlsx"
    workbook = load_workbook(report_path, read_only=True)
    try:
        assert "Comp Match Issues" in workbook.sheetnames
    finally:
        workbook.close()

    with zipfile.ZipFile(result.evidence_pack_path) as evidence_zip:
        assert "compensation_match_issues.csv" in evidence_zip.namelist()
