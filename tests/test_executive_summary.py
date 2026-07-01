import contextlib
import io
import json
import os
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from executive_summary import RECOMMENDED_ACTIONS, SEVERITY_BY_ISSUE_TYPE
from main import EngineConfig, run_prooflink_engine
from scripts.generate_realistic_demo import generate_dataset
from verify_proof import sha256_file, verify_evidence_pack_zip
from streamlit_app import (
    build_category_summary_rows,
    build_technical_evidence_summary,
    resolve_review_verification,
)


def _polished_plan_rules() -> dict:
    return {
        "eligibility_rule": "age21_and_1year",
        "service_days_required": 365,
        "age_required": 21,
        "align_first_month": False,
        "plan_match_config": {
            "match_formula_name": "50% up to 6% current-engine demo calibration",
            "match_type": "percent_of_comp",
            "match_rate": 0.50,
            "match_cap_pct": 0.06,
            "match_frequency": "per_payroll",
            "true_up_enabled": False,
            "true_up_enabled_column": "True-Up Enabled",
            "eligible_comp_columns": ["Plan Compensation"],
            "excluded_comp_columns": ["Fringe", "Reimbursement"],
            "employee_class_column": "Employee Class",
            "eligible_classes": ["Full-Time", "Part-Time"],
            "excluded_classes": ["Intern", "Union Excluded"],
            "absolute_tolerance": 5.00,
            "relative_tolerance_pct": 0.15,
        },
    }


def test_priority_and_recommended_action_mappings_are_centralized():
    assert SEVERITY_BY_ISSUE_TYPE["EMPLOYMENT_STATUS_CONFLICT"] == "High"
    assert SEVERITY_BY_ISSUE_TYPE["POST_TERMINATION_COMPENSATION"] == "High"
    assert SEVERITY_BY_ISSUE_TYPE["LATE_CONTRIBUTION"] == "Low"
    assert "Reconcile payroll and recordkeeper" in RECOMMENDED_ACTIONS["DEFERRAL_MISMATCH"]
    assert "Confirm current participant status" in RECOMMENDED_ACTIONS["EMPLOYMENT_STATUS_CONFLICT"]


def test_polished_fixture_executive_summary_integration_without_ground_truth(tmp_path: Path, monkeypatch):
    paths = generate_dataset(
        employees=100,
        year=2025,
        seed=20250630,
        output_dir=tmp_path / "input",
        periods=8,
        prefix="prooflink_polished_100",
        profile="polished",
    )
    # Prove runtime summary generation does not read the synthetic ground-truth file.
    paths["ground_truth"].unlink()
    monkeypatch.setenv("MAPPING_YAML_PATH", str(paths["mapping"]))

    with contextlib.redirect_stdout(io.StringIO()):
        result = run_prooflink_engine(
            payroll_path=str(paths["payroll"]),
            rk_path=str(paths["recordkeeper"]),
            config=EngineConfig(
                plan_name="ProofLink Polished Demo Plan",
                payroll_vendor_hint="ADP",
                rk_vendor_hint="VENDOR_RK_1",
                output_dir=str(tmp_path / "output"),
                proofs_dir=str(tmp_path / "proofs"),
            ),
            run_id="executive-summary-test",
            plan_rules=_polished_plan_rules(),
        )

    summary = result.summary["executive_summary"]
    json.dumps(summary)

    assert summary["run_id"] == "executive-summary-test"
    assert summary["plan_name"] == "ProofLink Polished Demo Plan"
    assert summary["plan_year"] == 2025
    assert summary["employees_reviewed"] == 100
    assert summary["payroll_periods_reviewed"] == 8
    assert summary["payroll_row_count"] == 800
    assert summary["recordkeeper_row_count"] == 793
    assert summary["total_exceptions"] == 12
    assert summary["high_priority_count"] == 4
    assert summary["medium_priority_count"] == 5
    assert summary["low_priority_count"] == 3
    assert summary["skipped_analyzer_warnings"] == []

    assert summary["counts_by_issue_category"] == {
        "Core Reconciliation": 5,
        "Contribution Timing": 2,
        "Population Validation": 2,
        "Compensation/Match": 2,
        "Secure 2.0": 1,
    }
    assert summary["counts_by_issue_type"]["Deferral amount mismatch"] == 2
    assert summary["counts_by_issue_type"]["Late contribution deposit"] == 2
    assert summary["counts_by_issue_type"]["Employment status conflict"] == 1
    assert summary["counts_by_issue_type"]["Compensation after termination"] == 1
    assert len(summary["recommended_actions"]) >= 8

    artifacts = result.summary["executive_summary_artifacts"]
    for key in ["executive_summary_json", "executive_summary_html", "all_exceptions_csv"]:
        assert Path(artifacts[key]).exists(), key

    all_exceptions = pd.read_csv(artifacts["all_exceptions_csv"])
    assert len(all_exceptions) == 12
    assert all_exceptions["issue_category"].value_counts().to_dict() == summary["counts_by_issue_category"]

    workbook = load_workbook(tmp_path / "output" / "reconciliation_report.xlsx", read_only=True)
    try:
        assert workbook.sheetnames[0] == "Review Summary"
        assert "Summary" in workbook.sheetnames
        assert "Population Issues" in workbook.sheetnames
    finally:
        workbook.close()

    zip_result = verify_evidence_pack_zip(Path(result.evidence_pack_path))
    assert zip_result["overall_ok"] is True

    with open(result.evidence_pack_path, "rb") as handle:
        assert handle.read(2) == b"PK"

    manifest_outputs = result.manifest.get("outputs", {})
    assert "executive_summary_json" in manifest_outputs
    assert "executive_summary_html" in manifest_outputs
    assert "all_exceptions_csv" in manifest_outputs
    assert "excel_report" in manifest_outputs

    for key in ["executive_summary_json", "executive_summary_html", "all_exceptions_csv"]:
        artifact_path = Path(artifacts[key])
        assert manifest_outputs[key]["sha256"] == sha256_file(artifact_path)
        assert Path(manifest_outputs[key]["path"]).name == artifact_path.name

    assert summary["verification_status"]["input_verification"] == "PASS"
    assert summary["verification_status"]["output_verification"] == "PASS"
    assert summary["verification_status"]["overall_verification"] == "PASS"

    display_verification = resolve_review_verification(result.summary, summary)
    assert display_verification["overall_verification"] == "PASS"

    category_rows = build_category_summary_rows(summary)
    assert category_rows == [
        {"Category": "Core Reconciliation", "Count": 5},
        {"Category": "Contribution Timing", "Count": 2},
        {"Category": "Compensation/Match", "Count": 2},
        {"Category": "Population Validation", "Count": 2},
        {"Category": "Secure 2.0", "Count": 1},
    ]
    assert all(isinstance(row["Count"], int) for row in category_rows)

    tech_summary = build_technical_evidence_summary(result.summary, summary, result.manifest, result.run_id, display_verification, evidence_zip_available=True)
    assert tech_summary["Verification status"] == "PASS"
    assert tech_summary["Output files verified"] >= 7
    assert tech_summary["Evidence ZIP available"] == "Yes"
    assert tech_summary["Manifest available"] == "Yes"
    assert not any(str(tmp_path) in str(value) for value in tech_summary.values())


def test_streamlit_review_summary_layout_uses_compact_status_and_technical_evidence():
    source = Path("streamlit_app.py").read_text(encoding="utf-8")
    assert "metric(\"Review status\"" not in source
    assert "**Review status:** Needs review" in source
    assert "build_category_summary_rows" in source
    assert "build_technical_evidence_summary" in source
    assert "Raw manifest and hashes" in source
    assert "st.json(summary.get(\"evidence_index\"" not in source
